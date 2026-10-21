# -*- coding: utf-8 -*-
"""Xuất fb/HUONG-DAN-TU-CHAY-QUANG-CAO.md thành file Word để gửi khách.
Chạy: python3 scripts/tao-huong-dan-qc-word.py"""
import os, re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "fb", "HUONG-DAN-TU-CHAY-QUANG-CAO.md")
OUT = os.path.join(ROOT, "fb", "Huong-dan-tu-chay-quang-cao.docx")

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hexcolor); tcPr.append(shd)

def add_runs(par, text):
    """**đậm** và `mã` trong một dòng."""
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part: continue
        if part.startswith("**"): r = par.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`"): r = par.add_run(part[1:-1]); r.font.name = "Consolas"
        else: par.add_run(part)

def build():
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    for s in doc.sections: s.left_margin = s.right_margin = Cm(2); s.top_margin = s.bottom_margin = Cm(1.8)
    lines = open(SRC, encoding="utf-8").read().splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("# "): doc.add_heading(ln[2:], 0)
        elif ln.startswith("## "): doc.add_heading(ln[3:], 1)
        elif ln.startswith("### "): doc.add_heading(ln[4:], 2)
        elif ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.match(r"^\|[\s\-|]+\|$", lines[i]): rows.append([c.strip() for c in lines[i].strip("|").split("|")])
                i += 1
            t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, val in enumerate(row):
                    cell = t.cell(r, c); cell.text = ""; add_runs(cell.paragraphs[0], val)
                    if r == 0: shade(cell, "E8F1F5"); cell.paragraphs[0].runs[0].bold = True
            doc.add_paragraph(); continue
        elif re.match(r"^\d+\. ", ln): add_runs(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\. ", "", ln))
        elif ln.startswith("- "): add_runs(doc.add_paragraph(style="List Bullet"), ln[2:])
        elif ln.strip() == "---": doc.add_page_break()
        elif ln.strip(): add_runs(doc.add_paragraph(), ln)
        i += 1
    doc.save(OUT); print("ok", OUT)

if __name__ == "__main__": build()
